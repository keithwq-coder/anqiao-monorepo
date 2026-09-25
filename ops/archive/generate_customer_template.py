# Python Script to Generate Customer Import Template

"""
Generate Excel template for Feishu CRM import.
Copy this data into Excel and then import to your Feishu multitable.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def create_customer_template():
    """Create customer information template for Feishu CRM"""
    wb = Workbook()
    
    # Create customer info sheet
    ws1 = wb.active
    ws1.title = "客户信息表"
    
    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    # Headers based on the recommended fields
    headers = [
        "客户编号", "客户名称", "联系人", "联系电话", "行业分类", 
        "客户等级", "跟进状态", "商关联金额 (万元)", 
        "预计成交时间", "负责人", "备注", "下次跟进时间", "商机来源"
    ]
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    # Sample data (you can replace with your actual customer data)
    sample_data = [
        ["CL20260724001", "XX 银行", "张总", "13800138001", "银行", 
         "B 级", "洽谈中", 50, "2026-09-15", "张三", "预算有限，等待资金到位", 
         "2026-08-05", "新增"],
        
        ["CL20260724002", "中国电信分公司", "李经理", "13900139001", "电信", 
         "A 级", "报价中", 35, "2026-08-30", "张三", "技术方案已确认", 
         "2026-07-30", "新增"],
        
        ["CL20260724003", "XX 保险公司", "王总", "13700137001", "保险", 
         "B 级", "方案演示", 25, "2026-10-15", "李四", "刚完成首次演示", 
         "2026-07-31", "新增"],
        
        ["CL20260724004", "YY 商业银行", "赵行长", "13500135001", "银行", 
         "A 级", "未接触", 80, "2026-12-31", "李四", "新开发客户", 
         "2026-08-01", "新增"],
        
        ["CL20260724005", "ZZ 财险公司", "周总", "13300133001", "保险", 
         "C 级", "初次接触", 10, "2026-11-30", "王五", "初步沟通阶段", 
         "2026-08-10", "新增"],
    ]
    
    # Write sample data
    for row_num, data in enumerate(sample_data, 2):
        for col_num, value in enumerate(data, 1):
            cell = ws1.cell(row=row_num, column=col_num, value=value)
    
    # Save file
    output_file = "飞书_CRM_客户导入模板.xlsx"
    wb.save(output_file)
    print(f"✅ Template generated: {output_file}")
    print("\nUsage instructions:")
    print("1. Open the Excel file")
    print("2. Replace sample data with your actual customer information")
    print("3. Upload to Feishu via: Workbench → Multi-table → Import")
    return output_file

if __name__ == "__main__":
    create_customer_template()