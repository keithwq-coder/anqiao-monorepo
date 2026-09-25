import os, sys, re
import pdfplumber
from pptx import Presentation
import openpyxl
import docx

BASE = "C:/Users/K/Documents/安樵"
OUT = "D:/Project/中科安樵/wiki/.reasonix/extract"
os.makedirs(OUT, exist_ok=True)

# (label, absolute path)
FILES = [
    ("M01_company_intro", f"{BASE}/工作台/01-营销/公司介绍/2026中科安樵-公司介绍-产品介绍.pdf"),
    ("M02_laobohui", f"{BASE}/工作台/02-市场与渠道/行业活动/2026广州老博会邀请函.pdf"),
    ("M03_prod_solution_en", f"{BASE}/中科安樵产品与解决方案2026(英文版).pptx"),
    ("M03_catalogue", f"{BASE}/中科安樵产品目录_中英双语_2026.docx"),
    ("M03_manual_v5", f"{BASE}/工作台/01-营销/产品说明/中科安樵_产品说明书_v5_20260729.pptx"),
    ("M04_manual_v5", f"{BASE}/工作台/01-营销/产品说明/中科安樵_产品说明书_v5_20260729.pptx"),
    ("M04_core_adv", f"{BASE}/工作台/01-营销/中科安樵AI健康守护仪_产品核心优势_0804.docx"),
    ("M04_manual_v2", f"{BASE}/工作台/04-技术资料/中科安樵_产品说明书_v2.pdf"),
    ("M04_devicelist", f"{BASE}/工作台/01-营销/产品说明/产品参数-设备型号清单.xlsx"),
    ("M05_longhu", f"{BASE}/工作台/01-营销/长护险基金监管智能感知设备应用方案.pptx"),
    ("M05_kaijian", f"{BASE}/工作台/01-营销/方案提案/凯健_安樵_智能康养合作方案-V2.pptx"),
    ("M05_guoshou", f"{BASE}/工作台/01-营销/方案提案/国寿嘉园·雅境智慧康养社区建设实施方案v5.pptx"),
    ("M06_dealer_v3_pdf", f"{BASE}/工作台/02-市场与渠道/渠道政策/中科安樵经销商合作手册_v3.pdf"),
    ("M06_role_compare", f"{BASE}/工作台/02-市场与渠道/渠道政策/渠道角色对比表.pdf"),
    ("M06_property", f"{BASE}/工作台/02-市场与渠道/渠道政策/物业合作·业主沟通手册.pdf"),
    ("M07_price0820", f"{BASE}/工作台/03-协议合同/价格报价/价格表0820.xlsx"),
    ("M07_partner0820", f"{BASE}/工作台/03-协议合同/价格报价/合伙人政策0820.xlsx"),
    ("M07_retail0820", f"{BASE}/工作台/03-协议合同/价格报价/产品零售报价单_0820.docx"),
    ("M08_quality", f"{BASE}/工作台/05-行政/制度流程/中科安樵_产品质量保障政策_V1.0.docx"),
    ("M08_fulfill_sop", f"{BASE}/工作台/05-行政/制度流程/中科安樵_经销商订单履约SOP.docx"),
    ("M08_manual_v2", f"{BASE}/工作台/04-技术资料/中科安樵_产品说明书_v2.pdf"),
    ("M09_app_manual", f"{BASE}/工作台/04-技术资料/软件说明/APP使用手册（图文完整版）V2.docx"),
    ("M09_device_book", f"{BASE}/工作台/04-技术资料/产品手册/AI健康守护仪使用书完整版.docx"),
    ("M09_manual_v2", f"{BASE}/工作台/04-技术资料/中科安樵_产品说明书_v2.pdf"),
    ("M10_crm", f"{BASE}/工作台/02-市场与渠道/销售线索/中科安樵_飞书CRM模板_导入用.xlsx"),
    ("M11_dealer_agree", f"{BASE}/工作台/03-协议合同/合作协议/经销商合作协议书.docx"),
    ("M11_partner_agree", f"{BASE}/工作台/03-协议合同/合作协议/合伙人合作协议_模板.docx"),
    ("M12_longhu", f"{BASE}/工作台/01-营销/长护险基金监管智能感知设备应用方案.pptx"),
    ("M12_longhu_pdf", f"{BASE}/工作台/01-营销/长护险基金监管智能感知设备应用方案.pdf"),
    ("M12_guoshou", f"{BASE}/工作台/01-营销/方案提案/国寿嘉园·雅境智慧康养社区建设实施方案v5.pptx"),
    ("M12_kaijian", f"{BASE}/工作台/01-营销/方案提案/凯健_安樵_智能康养合作方案-V2.pptx"),
    ("M13_property", f"{BASE}/工作台/02-市场与渠道/渠道政策/物业合作·业主沟通手册.pdf"),
    ("M13_dealer_v3", f"{BASE}/工作台/02-市场与渠道/渠道政策/中科安樵经销商合作手册_v3.pdf"),
    ("M14_incentive", f"{BASE}/工作台/05-行政/人事/销售激励政策0820.docx"),
    ("M15_manual_v5", f"{BASE}/工作台/01-营销/产品说明/中科安樵_产品说明书_v5_20260729.pptx"),
    ("M16_channel_mgmt", f"{BASE}/工作台/05-行政/制度流程/中科安樵_渠道管理制度_草案V2.0.docx"),
    ("M16_incentive", f"{BASE}/工作台/05-行政/人事/销售激励政策0820.docx"),
]

def clean_cjk_space(text):
    # remove single spaces between CJK chars (letter-spacing artifacts)
    text = re.sub(r'(?<=[\u4e00-\u9fff])\s(?=[\u4e00-\u9fff])', '', text)
    return text

def extract_pdf(path):
    out=[]
    with pdfplumber.open(path) as pdf:
        for i,page in enumerate(pdf.pages):
            t=page.extract_text() or ""
            out.append(f"--- P{i+1} ---\n"+clean_cjk_space(t))
    return "\n".join(out)

def extract_pptx(path):
    out=[]
    prs=Presentation(path)
    for i,slide in enumerate(prs.slides):
        out.append(f"--- SLIDE {i+1} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line="".join(run.text for run in para.runs)
                    if not line:
                        line=shape.text_frame.text
                    if line.strip():
                        out.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(out)

def extract_docx(path):
    d=docx.Document(path)
    out=[]
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for ti,t in enumerate(d.tables):
        out.append(f"--- TABLE {ti+1} ---")
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)

def extract_xlsx(path):
    wb=openpyxl.load_workbook(path, data_only=True, read_only=True)
    out=[]
    for ws in wb.worksheets:
        out.append(f"=== SHEET: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            cells=[str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)

for label,path in FILES:
    if not os.path.exists(path):
        print(f"MISSING {label}: {path}")
        continue
    try:
        ext=path.lower()
        if ext.endswith(".pdf"):
            txt=extract_pdf(path)
        elif ext.endswith(".pptx"):
            txt=extract_pptx(path)
        elif ext.endswith(".docx"):
            txt=extract_docx(path)
        elif ext.endswith(".xlsx"):
            txt=extract_xlsx(path)
        else:
            txt=""
        with open(os.path.join(OUT,f"{label}.txt"),"w",encoding="utf-8") as f:
            f.write(txt)
        print(f"OK {label}: {len(txt)} chars -> {label}.txt")
    except Exception as e:
        print(f"ERR {label}: {e}")
